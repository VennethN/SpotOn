"""Rough renderer for a pptxgenjs deck: draws backgrounds, shapes, pictures and wrapped text
with DejaVu Sans (wider than Calibri, so any text that fits here fits in PowerPoint)."""
import sys, io
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

SCALE = 144  # px per inch
REG = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
fonts = {}
def font(size_pt, bold):
    key = (round(size_pt * 2), bold)
    if key not in fonts:
        fonts[key] = ImageFont.truetype(BOLD if bold else REG, max(4, int(size_pt / 72 * SCALE * 0.92)))
    return fonts[key]

def px(emu): return int(Emu(emu).inches * SCALE)

def colour(fill):
    try:
        if fill.type is None: return None
        rgb = fill.fore_color.rgb
        return '#' + str(rgb)
    except Exception:
        return None

def wrap(draw, text, f, width):
    lines = []
    for para in text.split('\n'):
        words = para.split(' '); line = ''
        for w in words:
            t = (line + ' ' + w).strip()
            if draw.textlength(t, font=f) <= width or not line: line = t
            else: lines.append(line); line = w
        lines.append(line)
    return lines

def render(pptx, out_prefix):
    prs = Presentation(pptx)
    W, H = px(prs.slide_width), px(prs.slide_height)
    for i, slide in enumerate(prs.slides, 1):
        bg = '#FFFFFF'
        try:
            if slide.background.fill.type is not None: bg = '#' + str(slide.background.fill.fore_color.rgb)
        except Exception: pass
        im = Image.new('RGBA', (W, H), bg); draw = ImageDraw.Draw(im)
        overflow = []
        for sh in slide.shapes:
            x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pic = Image.open(io.BytesIO(sh.image.blob)).convert('RGBA')
                    crop = sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom
                    if any(c for c in crop):
                        pw, ph = pic.size
                        pic = pic.crop((int(pw*crop[0]), int(ph*crop[1]), int(pw*(1-crop[2])), int(ph*(1-crop[3]))))
                    pic = pic.resize((max(1, w), max(1, h)))
                    im.alpha_composite(pic, (x, y))
                except Exception as e: draw.rectangle([x, y, x+w, y+h], outline='red')
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                c = colour(sh.fill)
                if c: draw.rounded_rectangle([x, y, x+w, y+h], radius=int(min(w, h)*0.12) if 'Rounded' in str(sh.auto_shape_type) or 'Hex' in str(sh.auto_shape_type) else 0, fill=c)
            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                cy = y
                total_lines = []
                for p in tf.paragraphs:
                    runs = p.runs
                    if not runs: total_lines.append(('', None, 0)); continue
                    size = None; bold = False; col = '#000000'
                    for r in runs:
                        if r.font.size: size = r.font.size.pt
                        if r.font.bold: bold = True
                        try:
                            if r.font.color and r.font.color.rgb: col = '#' + str(r.font.color.rgb)
                        except Exception: pass
                    size = size or 18
                    f = font(size, bold)
                    text = ''.join(r.text for r in runs)
                    bullet = p._p.pPr is not None and p._p.pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None
                    indent = int(0.25 * SCALE) if bullet else 0
                    for ln in wrap(draw, text, f, w - indent - 4):
                        total_lines.append((ln, f, size, col, indent, bullet))
                    if p.space_after: total_lines.append(('', None, p.space_after.pt * 0.5))
                # vertical alignment
                heights = [(l[2] / 72 * SCALE * 1.2 if l[1] else l[2] / 72 * SCALE) for l in total_lines]
                th = sum(heights)
                anchor = tf._txBody.bodyPr.get('anchor')
                if anchor == 'ctr': cy = y + (h - th) / 2
                elif anchor == 'b': cy = y + h - th
                align = tf.paragraphs[0].alignment
                for l, lh in zip(total_lines, heights):
                    if l[1]:
                        ln, f, size, col, indent, bullet = l
                        tw = draw.textlength(ln, font=f)
                        lx = x + indent
                        if str(align).startswith('CENTER'): lx = x + (w - tw) / 2
                        if bullet and indent: draw.text((x + 4, cy), '•', font=f, fill=col)
                        draw.text((lx, cy), ln, font=f, fill=col)
                    cy += lh
                if th > h + 2: overflow.append(f"'{tf.text[:40]}' needs {th/SCALE:.2f}in, box {h/SCALE:.2f}in")
        im.convert('RGB').save(f'{out_prefix}-{i}.png')
        print(f'slide {i}:', 'ok' if not overflow else 'OVERFLOW ' + ' | '.join(overflow))

render(sys.argv[1], sys.argv[2])
